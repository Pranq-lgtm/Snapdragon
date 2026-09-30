from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt
from reportlab.lib import colors
from reportlab.lib.pagesizes import landscape, letter
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "presentation"
OUT.mkdir(exist_ok=True)

SLIDES = [
    ("Snapdragon Local Review", ["Private, offline-first code and Vercel preflight", "NPU-accelerated developer feedback for Snapdragon PCs"]),
    ("The problem", ["Cloud code review creates latency and data-governance risk", "Manual pre-deployment checks miss secrets and configuration errors", "Local builds already consume the CPU developers need for feedback"]),
    ("The solution", ["A single local CLI before commit or deployment", "Deterministic checks catch actionable issues immediately", "An optional local coding model explains and remediates findings", "No API keys, no source upload, no cloud fallback"]),
    ("Technical architecture", ["Repository → CLI → deterministic rules", "Optional adapter → tokenizer → ONNX Runtime", "QNN Execution Provider → Hexagon NPU / HTP", "CPU remains available for compilers, Docker, tests, and local servers"]),
    ("Why Snapdragon", ["Privacy: proprietary code remains on-device", "Responsiveness: inference is offloaded from the CPU", "Offline operation: useful on flights and restricted networks", "Enterprise fit: local governance and predictable data boundaries"]),
    ("Developer workflow", ["snapdragon-review check .", "JSON output integrates with editors and CI", "--strict turns warnings into deployment blockers", "Pre-commit catches issues before staging or Vercel"]),
    ("Demo and measurement", ["Inject a token and insecure Vercel rewrite", "Run the CLI and show structured findings", "Show QNN provider availability on ARM64", "Compare review latency and CPU utilization during a build"]),
    ("Accessibility and roadmap", ["Works today without AI dependencies", "Add Transformers CPU fallback for compatibility", "Add ONNX Runtime GenAI/QNN adapter for NPU inference", "Expand rules, VS Code integration, and signed model bundles"]),
    ("Call to action", ["Install locally", "Keep the review loop private", "Turn every Snapdragon developer machine into an edge code-review station"]),
]


def add_text_box(slide, text, x, y, w, h, size=24, color=(20, 32, 60), bold=False):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    p = frame.paragraphs[0]
    p.text = text
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = RGBColor(*color)
    return box


def make_pptx() -> None:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    for index, (title, bullets) in enumerate(SLIDES):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        bg = slide.background.fill
        bg.solid()
        bg.fore_color.rgb = RGBColor(246, 248, 252)
        add_text_box(slide, "SNAPDRAGON LOCAL REVIEW", 0.7, 0.45, 7, 0.3, 12, (30, 110, 210), True)
        add_text_box(slide, title, 0.7, 1.05, 11.8, 0.8, 34, (20, 32, 60), True)
        for line, bullet in enumerate(bullets):
            add_text_box(slide, f"• {bullet}", 1.0, 2.15 + line * 0.75, 11.3, 0.55, 22, (45, 55, 75))
        add_text_box(slide, f"{index + 1:02d}", 12.2, 6.75, 0.5, 0.3, 12, (100, 110, 130), True)
    prs.save(OUT / "snapdragon-local-review.pptx")


def make_pdf() -> None:
    styles = getSampleStyleSheet()
    title = styles["Title"]
    title.textColor = colors.HexColor("#14203c")
    body = styles["BodyText"]
    body.fontSize = 18
    body.leading = 25
    doc = SimpleDocTemplate(str(OUT / "snapdragon-local-review.pdf"), pagesize=landscape(letter), rightMargin=54, leftMargin=54, topMargin=42, bottomMargin=42)
    story = []
    for index, (heading, bullets) in enumerate(SLIDES):
        story.append(Paragraph(heading, title))
        story.append(Spacer(1, 24))
        for bullet in bullets:
            story.append(Paragraph(f"• {bullet}", body))
            story.append(Spacer(1, 12))
        story.append(Spacer(1, 24))
        story.append(Paragraph(f"{index + 1:02d}  |  SNAPDRAGON LOCAL REVIEW", styles["Caption"]))
        if index != len(SLIDES) - 1:
            story.append(Spacer(1, 280))
    doc.build(story)


if __name__ == "__main__":
    make_pptx()
    make_pdf()
